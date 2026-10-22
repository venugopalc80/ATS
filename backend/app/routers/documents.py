import hashlib
import io
import logging
import os
import re
from pathlib import Path
from datetime import datetime
from uuid import UUID, uuid4

from fastapi import APIRouter, BackgroundTasks, Depends, File, HTTPException, Query, UploadFile
from docx import Document as DocxDocument
from pypdf import PdfReader
from pydantic import BaseModel
from supabase import create_client

from app.auth import assert_org_member, get_current_user_id
from app.db import DatabaseConfigurationError, DatabaseConnectionError, get_connection

router = APIRouter(prefix="/api/documents", tags=["candidate documents"])
logger = logging.getLogger(__name__)
BUCKET = "candidate-documents"
MAX_BYTES = 5 * 1024 * 1024
ALLOWED = {
    "application/pdf": (".pdf", b"%PDF-"),
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": (".docx", b"PK\x03\x04"),
}


class DocumentOut(BaseModel):
    id: UUID
    organization_id: UUID
    candidate_id: UUID
    application_id: UUID | None = None
    document_type: str
    original_filename: str
    content_type: str
    file_size_bytes: int
    sha256: str
    created_at: datetime | None = None


def validate_document_bytes(data: bytes, mime: str) -> str:
    """Validate the allowlisted file type, bounded size, and basic file signature."""
    if mime not in ALLOWED:
        raise HTTPException(status_code=415, detail="Only PDF and DOCX files are supported")
    if not data:
        raise HTTPException(status_code=422, detail="The uploaded file is empty")
    if len(data) > MAX_BYTES:
        raise HTTPException(status_code=413, detail="CV files must be 5 MB or smaller")
    extension, signature = ALLOWED[mime]
    if not data.startswith(signature):
        raise HTTPException(status_code=415, detail="The file contents do not match the selected file type")
    return extension


def storage_client():
    url = os.getenv("SUPABASE_URL")
    key = os.getenv("SUPABASE_SERVICE_ROLE_KEY")
    if not url or not key:
        raise HTTPException(status_code=503, detail="Private document storage is not configured")
    return create_client(url, key)


def verify_candidate(candidate_id: UUID, organization_id: UUID):
    with get_connection() as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                "select id from public.candidates where id=%s and organization_id=%s",
                [candidate_id, organization_id],
            )
            if cursor.fetchone() is None:
                raise HTTPException(status_code=404, detail="Candidate not found")


def safe_filename(filename: str | None, extension: str) -> str:
    raw = Path(filename or "resume" + extension).name
    cleaned = re.sub(r"[^A-Za-z0-9._ -]", "_", raw).strip(" .")
    if not cleaned:
        cleaned = "resume" + extension
    if not cleaned.lower().endswith(extension):
        cleaned = Path(cleaned).stem + extension
    return cleaned[:180]


@router.post("", status_code=201, response_model=DocumentOut)
async def upload_candidate_document(
    organization_id: UUID,
    candidate_id: UUID,
    file: UploadFile = File(...),
    application_id: UUID | None = Query(default=None),
    user_id: UUID = Depends(get_current_user_id),
):
    assert_org_member(user_id, organization_id)
    verify_candidate(candidate_id, organization_id)

    mime = (file.content_type or "").lower()
    data = await file.read(MAX_BYTES + 1)
    extension = validate_document_bytes(data, mime)

    if application_id:
        with get_connection() as connection:
            with connection.cursor() as cursor:
                cursor.execute(
                    """select id from public.applications
                       where id=%s and organization_id=%s and candidate_id=%s""",
                    [application_id, organization_id, candidate_id],
                )
                if cursor.fetchone() is None:
                    raise HTTPException(status_code=404, detail="Application not found for this candidate")

    document_id = uuid4()
    path = f"{organization_id}/{candidate_id}/{document_id}{extension}"
    filename = safe_filename(file.filename, extension)
    digest = hashlib.sha256(data).hexdigest()

    try:
        storage_client().storage.from_(BUCKET).upload(
            path,
            data,
            {"content-type": mime, "upsert": "false", "cache-control": "0"},
        )
        with get_connection() as connection:
            with connection.cursor() as cursor:
                cursor.execute(
                    """insert into public.candidate_documents
                       (id, organization_id, candidate_id, application_id, document_type,
                        original_filename, storage_bucket, storage_path, content_type,
                        file_size_bytes, sha256, uploaded_by)
                       values (%s,%s,%s,%s,'cv',%s,%s,%s,%s,%s,%s,%s)
                       returning id, organization_id, candidate_id, application_id, document_type,
                                 original_filename, content_type, file_size_bytes, sha256, created_at""",
                    [document_id, organization_id, candidate_id, application_id, filename,
                     BUCKET, path, mime, len(data), digest, user_id],
                )
                row = cursor.fetchone()
                cursor.execute(
                    "update public.candidates set resume_path=%s, updated_at=now() where id=%s and organization_id=%s",
                    [path, candidate_id, organization_id],
                )
                cursor.execute(
                    """insert into public.audit_events
                       (organization_id, actor_user_id, action, entity_type, entity_id, metadata)
                       values (%s,%s,'candidate_document_uploaded','candidate_document',%s,%s::jsonb)""",
                    [organization_id, user_id, document_id,
                     '{"candidate_id":"' + str(candidate_id) + '","document_type":"cv"}'],
                )
                return row
    except HTTPException:
        try:
            storage_client().storage.from_(BUCKET).remove([path])
        except Exception:
            pass
        raise
    except Exception as exc:
        try:
            storage_client().storage.from_(BUCKET).remove([path])
        except Exception:
            pass
        raise HTTPException(status_code=503, detail="Unable to store the document") from exc


@router.get("", response_model=list[DocumentOut])
async def list_candidate_documents(
    organization_id: UUID,
    candidate_id: UUID,
    user_id: UUID = Depends(get_current_user_id),
):
    assert_org_member(user_id, organization_id)
    try:
        verify_candidate(candidate_id, organization_id)
        with get_connection() as connection:
            with connection.cursor() as cursor:
                cursor.execute(
                    """select id, organization_id, candidate_id, application_id, document_type,
                              original_filename, content_type, file_size_bytes, sha256, created_at
                       from public.candidate_documents
                       where organization_id=%s and candidate_id=%s and deleted_at is null
                       order by created_at desc""",
                    [organization_id, candidate_id],
                )
                return list(cursor.fetchall())
    except HTTPException:
        raise
    except (DatabaseConfigurationError, DatabaseConnectionError) as exc:
        logger.exception("Database failure while listing candidate documents")
        raise HTTPException(status_code=503, detail="Document service unavailable; please retry shortly") from exc
    except Exception as exc:
        logger.exception("Unexpected failure while listing candidate documents")
        raise HTTPException(status_code=503, detail="Unable to load candidate documents; please retry shortly") from exc


@router.post("/{document_id}/download")
async def create_document_download(
    document_id: UUID,
    organization_id: UUID,
    user_id: UUID = Depends(get_current_user_id),
):
    assert_org_member(user_id, organization_id)
    try:
        with get_connection() as connection:
            with connection.cursor() as cursor:
                cursor.execute(
                    """select storage_bucket, storage_path, original_filename
                       from public.candidate_documents
                       where id=%s and organization_id=%s and deleted_at is null""",
                    [document_id, organization_id],
                )
                row = cursor.fetchone()
                if row is None:
                    raise HTTPException(status_code=404, detail="Document not found")
        result = storage_client().storage.from_(row["storage_bucket"]).create_signed_url(row["storage_path"], 60, options={"download": False})
        signed_url = result.get("signedURL") or result.get("signedUrl")
        if not signed_url:
            raise HTTPException(status_code=503, detail="Unable to create secure download link")
        return {"url": signed_url, "filename": row["original_filename"], "expires_in": 60}
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=503, detail="Document service unavailable") from exc


@router.post("/{document_id}/preview")
async def preview_candidate_document(
    document_id: UUID,
    organization_id: UUID,
    user_id: UUID = Depends(get_current_user_id),
):
    """Return a short-lived inline PDF URL or extracted DOCX text for an authorised recruiter."""
    assert_org_member(user_id, organization_id)
    try:
        with get_connection() as connection:
            with connection.cursor() as cursor:
                cursor.execute(
                    """select candidate_id, storage_bucket, storage_path, original_filename, content_type
                       from public.candidate_documents
                       where id=%s and organization_id=%s and deleted_at is null""",
                    [document_id, organization_id],
                )
                row = cursor.fetchone()
                if row is None:
                    raise HTTPException(status_code=404, detail="Document not found")
        if row["content_type"] == "application/pdf":
            result = storage_client().storage.from_(row["storage_bucket"]).create_signed_url(
                row["storage_path"], 60, options={"download": False}
            )
            signed_url = result.get("signedURL") or result.get("signedUrl")
            if not signed_url:
                raise HTTPException(status_code=503, detail="Unable to create secure preview link")
            return {"kind": "pdf", "url": signed_url, "filename": row["original_filename"], "expires_in": 60}

        if row["content_type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            data = storage_client().storage.from_(row["storage_bucket"]).download(row["storage_path"])
            text = extract_cv_text(data, row["content_type"])
            if not text.strip():
                raise HTTPException(status_code=422, detail="No readable text was found in this DOCX file")
            return {"kind": "text", "text": text, "filename": row["original_filename"]}

        raise HTTPException(status_code=415, detail="Preview is not supported for this file type")
    except HTTPException:
        raise
    except Exception as exc:
        logger.exception("Failed to preview candidate document")
        raise HTTPException(status_code=503, detail="Unable to preview this document right now") from exc


MAX_BULK_FILES = 20
MAX_BULK_BYTES = 4 * 1024 * 1024  # Stay below the Vercel Functions request-body limit.
CV_SKILL_KEYWORDS = (
    "python", "java", "javascript", "typescript", "react", "next.js", "node.js",
    "sql", "postgresql", "mysql", "mongodb", "aws", "azure", "gcp", "docker",
    "kubernetes", "machine learning", "deep learning", "tensorflow", "pytorch",
    "data analysis", "power bi", "tableau", "excel", "django", "fastapi",
    "project management", "salesforce", "account management", "recruitment",
)


def extract_cv_text(data: bytes, mime: str) -> str:
    """Extract selectable text from supported CV formats; scanned PDFs need OCR separately."""
    try:
        if mime == "application/pdf":
            reader = PdfReader(io.BytesIO(data))
            return "\n".join((page.extract_text() or "") for page in reader.pages[:20])[:100_000]
        if mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            document = DocxDocument(io.BytesIO(data))
            parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" ".join(cell.text for cell in row.cells))
            return "\n".join(parts)[:100_000]
    except Exception:
        return ""
    return ""


def cv_candidate_fields(text: str, filename: str) -> dict:
    email_match = re.search(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", text, re.I)
    phone_match = re.search(r"(?<!\w)(?:\+?\d[\d ()\-.]{7,}\d)(?!\w)", text)
    excluded = {
        "curriculum vitae", "resume", "résumé", "cv", "professional summary",
        "summary", "profile", "contact", "contact details", "work experience",
        "professional experience", "education", "skills", "experience",
    }
    name = ""
    for raw_line in text.splitlines()[:15]:
        line = re.sub(r"\s+", " ", raw_line).strip(" \t|•-")
        if not line or len(line) > 80 or "@" in line or re.search(r"\d{3}", line):
            continue
        if line.casefold() in excluded or len(line.split()) > 5:
            continue
        if re.fullmatch(r"[A-Za-zÀ-ÖØ-öø-ÿ'’. -]+", line) and len(line.split()) >= 2:
            name = line
            break
    if not name:
        stem = Path(filename).stem
        stem = re.sub(r"\b(cv|resume|résumé|curriculum vitae|final|updated|new)\b", " ", stem, flags=re.I)
        stem = re.sub(r"[_-]+", " ", stem)
        name = re.sub(r"\s+", " ", stem).strip(" .-_")
    parts = name.split()
    first_name = (parts[0] if parts else "CV candidate")[:100]
    last_name = (" ".join(parts[1:]) if len(parts) > 1 else None)
    skills = [skill for skill in CV_SKILL_KEYWORDS if re.search(r"(?<![a-z0-9])" + re.escape(skill) + r"(?![a-z0-9])", text, re.I)]
    return {
        "first_name": first_name,
        "last_name": last_name[:100] if last_name else None,
        "email": email_match.group(0).lower() if email_match else None,
        "phone": re.sub(r"\s+", " ", phone_match.group(0)).strip()[:50] if phone_match else None,
        "skills": skills[:30],
    }


@router.post("/bulk-cv", status_code=200)
async def bulk_upload_cvs(
    organization_id: UUID,
    files: list[UploadFile] = File(...),
    user_id: UUID = Depends(get_current_user_id),
):
    """Create candidate profiles from a batch of PDF/DOCX CVs and store each CV privately."""
    assert_org_member(user_id, organization_id)
    if not files:
        raise HTTPException(status_code=422, detail="Choose at least one CV to upload")
    if len(files) > MAX_BULK_FILES:
        raise HTTPException(status_code=413, detail=f"Upload up to {MAX_BULK_FILES} CVs per batch")

    declared_total = sum(file.size or 0 for file in files)
    if declared_total > MAX_BULK_BYTES:
        raise HTTPException(status_code=413, detail="The total batch must be 4 MB or smaller. Split larger batches into multiple uploads.")

    results = []
    for file in files:
        filename = Path(file.filename or "resume.pdf").name
        mime = (file.content_type or "").lower()
        data = await file.read(MAX_BYTES + 1)
        try:
            extension = validate_document_bytes(data, mime)
            filename = safe_filename(file.filename, extension)
            digest = hashlib.sha256(data).hexdigest()
            text = extract_cv_text(data, mime)
            fields = cv_candidate_fields(text, filename)
            if not fields["first_name"].strip():
                fields["first_name"] = "CV candidate"

            with get_connection() as connection:
                with connection.cursor() as cursor:
                    if fields["email"]:
                        cursor.execute(
                            "select id from public.candidates where organization_id=%s and lower(email)=lower(%s) limit 1",
                            [organization_id, fields["email"]],
                        )
                        existing = cursor.fetchone()
                        if existing:
                            results.append({"filename": filename, "status": "duplicate", "message": "A candidate with this email already exists", "candidate_id": str(existing["id"])})
                            continue
                    cursor.execute(
                        "select id from public.candidate_documents where organization_id=%s and sha256=%s and deleted_at is null limit 1",
                        [organization_id, digest],
                    )
                    if cursor.fetchone():
                        results.append({"filename": filename, "status": "duplicate", "message": "This CV file has already been uploaded"})
                        continue

            candidate_id = uuid4()
            document_id = uuid4()
            path = f"{organization_id}/{candidate_id}/{document_id}{extension}"
            storage_client().storage.from_(BUCKET).upload(
                path, data, {"content-type": mime, "upsert": "false", "cache-control": "0"}
            )
            try:
                with get_connection() as connection:
                    with connection.cursor() as cursor:
                        cursor.execute(
                            """insert into public.candidates
                               (id, organization_id, first_name, last_name, email, phone, status,
                                skills, resume_path, resume_text, created_at, updated_at)
                               values (%s,%s,%s,%s,%s,%s,'active',%s,%s,%s,now(),now())""",
                            [candidate_id, organization_id, fields["first_name"], fields["last_name"],
                             fields["email"], fields["phone"], fields["skills"], path, text or None],
                        )
                        cursor.execute(
                            """insert into public.candidate_documents
                               (id, organization_id, candidate_id, application_id, document_type,
                                original_filename, storage_bucket, storage_path, content_type,
                                file_size_bytes, sha256, uploaded_by)
                               values (%s,%s,%s,null,'cv',%s,%s,%s,%s,%s,%s,%s)""",
                            [document_id, organization_id, candidate_id, filename, BUCKET, path,
                             mime, len(data), digest, user_id],
                        )
                        cursor.execute(
                            """insert into public.audit_events
                               (organization_id, actor_user_id, action, entity_type, entity_id, metadata)
                               values (%s,%s,'candidate_cv_bulk_imported','candidate',%s,%s::jsonb)""",
                            [organization_id, user_id, candidate_id,
                             '{"filename":' + __import__("json").dumps(filename) + ',"document_id":"' + str(document_id) + '"}'],
                        )
                results.append({
                    "filename": filename, "status": "imported", "candidate_id": str(candidate_id),
                    "name": " ".join(part for part in [fields["first_name"], fields["last_name"]] if part),
                    "email": fields["email"], "skills_found": fields["skills"],
                    "text_extracted": bool(text.strip()),
                })
            except Exception:
                try:
                    storage_client().storage.from_(BUCKET).remove([path])
                except Exception:
                    pass
                raise
        except HTTPException as exc:
            results.append({"filename": filename, "status": "error", "message": str(exc.detail)})
        except Exception:
            results.append({"filename": filename, "status": "error", "message": "Could not import this CV. Check storage/database configuration and try again."})
        finally:
            await file.close()

    return {"total": len(files), "imported": sum(item["status"] == "imported" for item in results),
            "duplicates": sum(item["status"] == "duplicate" for item in results),
            "failed": sum(item["status"] == "error" for item in results), "results": results}


# Idea 06 — extraction is stored separately from confirmed candidate data.
PARSER_VERSION = "aqerith-cv-parser-v1"
CONFIRMABLE_FIELDS = {"first_name", "last_name", "email", "phone", "skills", "current_title", "years_experience"}


class ConfirmExtractionIn(BaseModel):
    fields: dict


def _document_for_extraction(document_id: UUID, organization_id: UUID) -> dict:
    with get_connection() as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """select id, organization_id, candidate_id, storage_bucket, storage_path,
                          original_filename, content_type
                   from public.candidate_documents
                   where id=%s and organization_id=%s and deleted_at is null""",
                [document_id, organization_id],
            )
            row = cursor.fetchone()
            if row is None:
                raise HTTPException(status_code=404, detail="Document not found")
            return row


def _structured_cv_fields(text: str, filename: str) -> tuple[dict, dict, dict]:
    fields = cv_candidate_fields(text, filename)
    lines = [re.sub(r"\\s+", " ", line).strip(" \\t|•-") for line in text.splitlines()]
    lines = [line for line in lines if line]
    sections = {"work_history": [], "education": [], "qualifications": []}
    current = None
    headings = {
        "work_history": {"experience", "work experience", "employment history", "professional experience", "career history"},
        "education": {"education", "academic background", "education history"},
        "qualifications": {"qualifications", "certifications", "certificates", "licenses", "licences"},
    }
    for line in lines:
        normalized = re.sub(r"[^a-z ]", "", line.casefold()).strip()
        matched = next((key for key, values in headings.items() if normalized in values), None)
        if matched:
            current = matched
            continue
        if normalized in {"skills", "technical skills", "summary", "professional summary", "profile", "projects", "references"}:
            current = None
            continue
        if current and len(line) <= 240:
            sections[current].append(line)
    data = {
        **fields,
        "work_history": sections["work_history"][:80],
        "education": sections["education"][:40],
        "qualifications": sections["qualifications"][:40],
    }
    confidence = {
        "first_name": 0.55 if fields.get("first_name") else 0.0,
        "last_name": 0.55 if fields.get("last_name") else 0.0,
        "email": 0.99 if fields.get("email") else 0.0,
        "phone": 0.75 if fields.get("phone") else 0.0,
        "skills": 0.65 if fields.get("skills") else 0.0,
        "work_history": 0.45 if sections["work_history"] else 0.0,
        "education": 0.45 if sections["education"] else 0.0,
        "qualifications": 0.4 if sections["qualifications"] else 0.0,
    }
    provenance = {
        "source": "original_document_text",
        "filename": filename,
        "method": "rule_based",
        "parser_version": PARSER_VERSION,
        "text_characters": len(text),
        "limitations": ["Scanned PDFs require OCR", "Confidence values are heuristic and require recruiter review"],
    }
    return data, confidence, provenance


def _run_document_extraction(document_id: UUID, organization_id: UUID, user_id: UUID) -> None:
    try:
        doc = _document_for_extraction(document_id, organization_id)
        with get_connection() as connection:
            with connection.cursor() as cursor:
                cursor.execute(
                    """update public.document_extractions
                       set status='processing', attempt_count=attempt_count+1,
                           started_at=now(), updated_at=now(), errors='[]'::jsonb
                       where organization_id=%s and document_id=%s and status <> 'confirmed'""",
                    [organization_id, document_id],
                )
        raw = storage_client().storage.from_(doc["storage_bucket"]).download(doc["storage_path"])
        text = extract_cv_text(raw, doc["content_type"])
        if not text.strip():
            raise ValueError("No readable text was found. The PDF may be scanned, corrupted or password-protected.")
        data, confidence, provenance = _structured_cv_fields(text, doc["original_filename"])
        import json
        with get_connection() as connection:
            with connection.cursor() as cursor:
                cursor.execute(
                    """update public.document_extractions
                       set status='review_required', parser_version=%s, extraction_data=%s::jsonb,
                           field_confidence=%s::jsonb, provenance=%s::jsonb, errors='[]'::jsonb,
                           completed_at=now(), updated_at=now()
                       where organization_id=%s and document_id=%s and status <> 'confirmed'""",
                    [PARSER_VERSION, json.dumps(data), json.dumps(confidence), json.dumps(provenance),
                     organization_id, document_id],
                )
                cursor.execute(
                    """insert into public.audit_events
                       (organization_id, actor_user_id, action, entity_type, entity_id, metadata)
                       values (%s,%s,'cv_extraction_completed','candidate_document',%s,%s::jsonb)""",
                    [organization_id, user_id, document_id,
                     json.dumps({"parser_version": PARSER_VERSION, "text_characters": len(text)})],
                )
    except Exception as exc:
        logger.exception("CV extraction failed for document %s", document_id)
        import json
        safe_error = str(exc)[:500] if isinstance(exc, (ValueError, HTTPException)) else "Document processing failed. Check file integrity and processing configuration."
        try:
            with get_connection() as connection:
                with connection.cursor() as cursor:
                    cursor.execute(
                        """update public.document_extractions
                           set status='failed', errors=%s::jsonb, completed_at=now(), updated_at=now()
                           where organization_id=%s and document_id=%s and status <> 'confirmed'""",
                        [json.dumps([safe_error]), organization_id, document_id],
                    )
        except Exception:
            logger.exception("Unable to persist CV extraction failure for document %s", document_id)


@router.post("/{document_id}/parse", status_code=202)
async def parse_candidate_document(
    document_id: UUID,
    background_tasks: BackgroundTasks,
    organization_id: UUID,
    user_id: UUID = Depends(get_current_user_id),
):
    assert_org_member(user_id, organization_id)
    doc = _document_for_extraction(document_id, organization_id)
    import json
    with get_connection() as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """insert into public.document_extractions
                   (organization_id, document_id, candidate_id, status, parser_version, requested_by)
                   values (%s,%s,%s,'queued',%s,%s)
                   on conflict (organization_id, document_id) do update
                     set status=case when public.document_extractions.status='confirmed'
                                     then 'confirmed' else 'queued' end,
                         requested_by=excluded.requested_by, updated_at=now(),
                         errors=case when public.document_extractions.status='confirmed'
                                     then public.document_extractions.errors else '[]'::jsonb end
                   returning id, status, attempt_count""",
                [organization_id, document_id, doc["candidate_id"], PARSER_VERSION, user_id],
            )
            state = cursor.fetchone()
            if state["status"] != "confirmed":
                cursor.execute(
                    """insert into public.audit_events
                       (organization_id, actor_user_id, action, entity_type, entity_id, metadata)
                       values (%s,%s,'cv_extraction_requested','candidate_document',%s,%s::jsonb)""",
                    [organization_id, user_id, document_id,
                     json.dumps({"parser_version": PARSER_VERSION})],
                )
    if state["status"] != "confirmed":
        background_tasks.add_task(_run_document_extraction, document_id, organization_id, user_id)
    return {"document_id": str(document_id), "status": state["status"], "processing": state["status"] != "confirmed"}


@router.get("/{document_id}/extraction")
async def get_document_extraction(
    document_id: UUID,
    organization_id: UUID,
    user_id: UUID = Depends(get_current_user_id),
):
    assert_org_member(user_id, organization_id)
    _document_for_extraction(document_id, organization_id)
    with get_connection() as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """select id, document_id, candidate_id, status, parser_version, attempt_count,
                          extraction_data, field_confidence, provenance, errors, requested_by,
                          started_at, completed_at, confirmed_at, confirmed_by, created_at, updated_at
                   from public.document_extractions
                   where organization_id=%s and document_id=%s""",
                [organization_id, document_id],
            )
            row = cursor.fetchone()
            if row is None:
                raise HTTPException(status_code=404, detail="No extraction has been requested for this document")
            return row


@router.post("/{document_id}/confirm-extraction")
async def confirm_document_extraction(
    document_id: UUID,
    payload: ConfirmExtractionIn,
    organization_id: UUID,
    user_id: UUID = Depends(get_current_user_id),
):
    assert_org_member(user_id, organization_id)
    doc = _document_for_extraction(document_id, organization_id)
    invalid = set(payload.fields.keys()) - CONFIRMABLE_FIELDS
    if invalid:
        raise HTTPException(status_code=422, detail="Unsupported candidate fields: " + ", ".join(sorted(invalid)))
    fields = payload.fields
    if "first_name" in fields and (not isinstance(fields["first_name"], str) or not fields["first_name"].strip()):
        raise HTTPException(status_code=422, detail="First name cannot be empty")
    if "email" in fields and fields["email"] is not None:
        email = str(fields["email"]).strip().lower()
        if email and not re.fullmatch(r"[^\\s@]+@[^\\s@]+\\.[^\\s@]+", email):
            raise HTTPException(status_code=422, detail="Email address is invalid")
        fields["email"] = email or None
    if "skills" in fields and not isinstance(fields["skills"], list):
        raise HTTPException(status_code=422, detail="Skills must be a list")
    if "years_experience" in fields and fields["years_experience"] is not None:
        try:
            years = float(fields["years_experience"])
        except (TypeError, ValueError) as exc:
            raise HTTPException(status_code=422, detail="Years of experience must be numeric") from exc
        if years < 0 or years > 80:
            raise HTTPException(status_code=422, detail="Years of experience must be between 0 and 80")
        fields["years_experience"] = years
    import json
    with get_connection() as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """select status, extraction_data from public.document_extractions
                   where organization_id=%s and document_id=%s for update""",
                [organization_id, document_id],
            )
            extraction = cursor.fetchone()
            if extraction is None or extraction["status"] not in {"review_required", "confirmed"}:
                raise HTTPException(status_code=409, detail="Extraction must finish successfully before confirmation")
            assignments = []
            values = []
            for key, value in fields.items():
                assignments.append(f"{key}=%s")
                values.append(value)
            if assignments:
                values.extend([doc["candidate_id"], organization_id])
                cursor.execute(
                    "update public.candidates set " + ", ".join(assignments) +
                    ", updated_at=now() where id=%s and organization_id=%s",
                    values,
                )
            confirmed_data = dict(extraction["extraction_data"] or {})
            confirmed_data.update(fields)
            cursor.execute(
                """update public.document_extractions
                   set status='confirmed', extraction_data=%s::jsonb, confirmed_at=now(),
                       confirmed_by=%s, updated_at=now()
                   where organization_id=%s and document_id=%s""",
                [json.dumps(confirmed_data), user_id, organization_id, document_id],
            )
            cursor.execute(
                """insert into public.audit_events
                   (organization_id, actor_user_id, action, entity_type, entity_id, metadata)
                   values (%s,%s,'cv_extraction_confirmed','candidate_document',%s,%s::jsonb)""",
                [organization_id, user_id, document_id,
                 json.dumps({"candidate_id": str(doc["candidate_id"]), "confirmed_fields": sorted(fields.keys())})],
            )
    return {"document_id": str(document_id), "candidate_id": str(doc["candidate_id"]), "status": "confirmed"}
